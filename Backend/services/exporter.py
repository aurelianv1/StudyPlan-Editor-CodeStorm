from docx import Document

from services.docx_utils import add_placeholder, document_to_bytes

EMPTY_SECTION_PLACEHOLDER = "[ Completați această secțiune ]"


def generate_docx(sections: list[dict]) -> bytes:
    """Generează un fișier .docx cu câte un titlu și un paragraf pentru fiecare secțiune."""
    doc = Document()

    for section in sections:
        doc.add_heading(section.get("title", ""), level=1)

        content = section.get("content", "")
        if content:
            doc.add_paragraph(content)
        else:
            add_placeholder(doc, EMPTY_SECTION_PLACEHOLDER)

    return document_to_bytes(doc)
