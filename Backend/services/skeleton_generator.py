from docx import Document
from docx.document import Document as DocxDocument

from services.docx_utils import add_placeholder, document_to_bytes

ACADEMIC_YEAR = "2025-2026"
FILL_IN = "[ Completați ]"
DEFAULT_EVALUATION_TYPE = "examen"
MAX_PROFESSIONAL_COMPETENCES = 5
MAX_TRANSVERSAL_COMPETENCES = 3


def generate_skeleton(plan_data: dict) -> bytes:
    """Generează un schelet de FD pre-completat din datele Planului de Învățământ."""
    doc = Document()
    evaluation_type = plan_data.get("tip_evaluare")
    credits = plan_data.get("credite")

    doc.add_heading("FIȘA DISCIPLINEI", level=0)

    doc.add_heading("1. Date despre programul de studiu", level=1)
    _add_table(
        doc,
        [
            ("Tip evaluare", evaluation_type.capitalize() if evaluation_type else FILL_IN),
            ("Număr credite", str(credits) if credits else FILL_IN),
            ("Anul universitar", ACADEMIC_YEAR),
        ],
    )
    doc.add_paragraph()

    doc.add_heading("2. Obiectivele disciplinei", level=1)
    add_placeholder(doc, "[ Completați obiectivele disciplinei ]")
    doc.add_paragraph()

    doc.add_heading("3. Competențe profesionale", level=1)
    _add_competences(
        doc,
        plan_data.get("competente_profesionale", []),
        prefix="CP",
        limit=MAX_PROFESSIONAL_COMPETENCES,
        placeholder="[ Completați competențele profesionale ]",
    )
    doc.add_paragraph()

    doc.add_heading("4. Competențe transversale", level=1)
    _add_competences(
        doc,
        plan_data.get("competente_transversale", []),
        prefix="CT",
        limit=MAX_TRANSVERSAL_COMPETENCES,
        placeholder="[ Completați competențele transversale ]",
    )
    doc.add_paragraph()

    doc.add_heading("5. Conținutul disciplinei", level=1)
    doc.add_heading("5.1 Curs", level=2)
    add_placeholder(doc, "[ Completați tematica cursului ]")
    doc.add_heading("5.2 Seminar/Laborator", level=2)
    add_placeholder(doc, "[ Completați tematica seminarului/laboratorului ]")
    doc.add_paragraph()

    doc.add_heading("6. Evaluare", level=1)
    _add_table(
        doc,
        [
            ("Tip activitate", "Pondere"),
            ((evaluation_type or DEFAULT_EVALUATION_TYPE).capitalize(), "60%"),
            ("Activitate laborator/seminar", "40%"),
        ],
    )
    doc.add_paragraph()

    doc.add_heading("7. Bibliografie", level=1)
    add_placeholder(
        doc, "[ Adăugați minim 3 surse bibliografice, dintre care una din ultimii 5 ani ]"
    )

    return document_to_bytes(doc)


def _add_table(doc: DocxDocument, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for row_index, (label, value) in enumerate(rows):
        table.cell(row_index, 0).text = label
        table.cell(row_index, 1).text = value


def _add_competences(
    doc: DocxDocument, competences: list[str], prefix: str, limit: int, placeholder: str
) -> None:
    relevant = [c for c in competences if c.startswith(prefix)][:limit]
    if not relevant:
        add_placeholder(doc, placeholder)
        return
    for competence in relevant:
        doc.add_paragraph(competence, style="List Bullet")
